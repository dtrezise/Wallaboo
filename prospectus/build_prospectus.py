from __future__ import annotations

import re
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


REPO = Path("/Users/dan/Documents/GitHub/Wallaboo")
SOURCE_MD = REPO / "prospectus" / "WALLABOO_PROSPECTUS_V0.1.md"
SOURCE_REGISTER = REPO / "research" / "SOURCE_REGISTER.md"
OUTPUT = REPO / "outputs" / "Wallaboo_Prospectus_v0.1.docx"
BUILD_DIR = Path("/private/tmp/wallaboo-prospectus-build")

NAVY = "15324B"
BLUE = "285B7E"
TEAL = "2A7F75"
GOLD = "D6A84B"
PALE_GOLD = "FFF4D6"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F2"
PALE_GRAY = "F4F6F9"
MID_GRAY = "D4DCE3"
DARK_GRAY = "44515C"
WHITE = "FFFFFF"
RED = "B42318"

PAGE_WIDTH_DXA = 12240
LEFT_MARGIN_DXA = 1440
RIGHT_MARGIN_DXA = 1440
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - LEFT_MARGIN_DXA - RIGHT_MARGIN_DXA


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int], indent=0) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        delta = CONTENT_WIDTH_DXA - sum(widths)
        widths[-1] += delta
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    for tag_name in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = tbl_pr.find(qn(tag_name))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_paragraph_control(paragraph, keep_next=False, keep_lines=False, widow=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for tag_name, enabled in (
        ("keepNext", keep_next),
        ("keepLines", keep_lines),
        ("widowControl", widow),
    ):
        old = p_pr.find(qn(f"w:{tag_name}"))
        if old is not None:
            p_pr.remove(old)
        if enabled:
            node = OxmlElement(f"w:{tag_name}")
            p_pr.append(node)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def ensure_list_abstract(doc: Document, kind: str) -> int:
    cache_name = f"_wallaboo_{kind}_abstract_num_id"
    cached = getattr(doc, cache_name, None)
    if cached is not None:
        return cached
    numbering = doc.part.numbering_part.element
    existing = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "260")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    setattr(doc, cache_name, abstract_id)
    return abstract_id


def new_list_instance(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_list_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08


def add_hyperlink(paragraph, text: str, url: str, color=BLUE) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def strip_md(text: str) -> str:
    return text.replace("**", "").replace("`", "").replace("*", "").strip()


def widths_for_table(headers: list[str]) -> list[int]:
    n = len(headers)
    if n == 2:
        return [2800, 6560]
    if n == 3:
        return [3100, 1700, 4560]
    if n == 4:
        return [2200, 1600, 1900, 3660]
    if n == 5:
        return [2440, 1000, 1100, 1500, 3320]
    if n == 6:
        return [2160, 1440, 1440, 1440, 1440, 1440]
    if n == 7:
        return [1450, 1320, 1320, 1320, 1320, 1320, 1310]
    return [CONTENT_WIDTH_DXA // n] * n


def looks_numeric(text: str) -> bool:
    value = strip_md(text)
    return bool(re.match(r"^[\$\(\)-]?[\d,.]+(?:\s*/\s*\d+)?(?:%|K|M)?\)?$", value))


def add_table(doc: Document, rows: list[list[str]], widths: list[int] | None = None, source_links=False):
    if not rows:
        return None
    cols = len(rows[0])
    rows = [r[:cols] + [""] * max(0, cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths or widths_for_table(rows[0]))
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    font_size = 7.6 if cols >= 6 else 8.2 if cols >= 5 else 8.7

    for ridx, row in enumerate(rows):
        set_cant_split(table.rows[ridx])
        is_total = ridx > 0 and strip_md(row[0]).lower() in {"total", "base"}
        for cidx, raw in enumerate(row):
            cell = table.cell(ridx, cidx)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if ridx == 0:
                set_cell_shading(cell, NAVY)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(strip_md(raw))
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
                run.font.size = Pt(font_size)
            else:
                if is_total:
                    set_cell_shading(cell, PALE_TEAL)
                if source_links and cidx == cols - 1:
                    urls = [u.strip() for u in raw.split(";") if u.strip()]
                    for uidx, url in enumerate(urls):
                        if uidx:
                            paragraph.add_run("  ")
                        add_hyperlink(paragraph, f"Source {uidx + 1}" if len(urls) > 1 else "Open source", url)
                else:
                    add_inline(paragraph, raw)
                if cidx > 0 and looks_numeric(raw):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    if is_total or "**" in raw or cidx == 0 and len(strip_md(raw)) < 25:
                        run.bold = True if is_total or "**" in raw else run.bold
            set_paragraph_control(paragraph, keep_lines=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        parts = [part.strip() for part in lines[idx].strip().strip("|").split("|")]
        if not all(re.match(r"^:?-{3,}:?$", part) for part in parts):
            rows.append(parts)
        idx += 1
    return rows, idx


def generate_charts() -> dict[str, Path]:
    if BUILD_DIR.exists():
        shutil.rmtree(BUILD_DIR)
    BUILD_DIR.mkdir(parents=True)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"

    def font(size: int, bold=False):
        path = bold_path if bold else regular_path
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            return ImageFont.load_default()

    def ink(value: str) -> str:
        return f"#{value}" if re.fullmatch(r"[0-9A-Fa-f]{6}", value) else value

    def centered(draw, xy, text, fnt, fill=NAVY):
        box = draw.textbbox((0, 0), text, font=fnt)
        draw.text((xy[0] - (box[2] - box[0]) / 2, xy[1]), text, font=fnt, fill=ink(fill))

    def wrapped_center(draw, x, y, lines, fnt, fill=DARK_GRAY, gap=4):
        for line in lines:
            centered(draw, (x, y), line, fnt, fill)
            y += fnt.size + gap

    market = BUILD_DIR / "market_context.png"
    img = Image.new("RGB", (1400, 650), "white")
    draw = ImageDraw.Draw(img)
    centered(draw, (700, 28), "Market context is large, but the definitions are not additive", font(34, True))
    left, right, top, bottom = 110, 1340, 105, 525
    max_value = 5.2
    for tick in range(0, 6):
        y = bottom - (tick / max_value) * (bottom - top)
        draw.line((left, y, right, y), fill="#D4DCE3", width=2)
        draw.text((25, y - 12), f"${tick}B", font=font(20), fill=ink(DARK_GRAY))
    values = [4.90, 3.66, 0.22]
    labels = [["U.S. Games/Puzzles", "2025"], ["U.S./Canada Hobby", "Games 2025"], ["Kickstarter Tabletop", "Pledges 2024"]]
    colors = [NAVY, TEAL, GOLD]
    centers = [350, 720, 1090]
    for value, label_lines, color, cx in zip(values, labels, colors, centers):
        height = (value / max_value) * (bottom - top)
        draw.rounded_rectangle((cx - 100, bottom - height, cx + 100, bottom), radius=8, fill=f"#{color}")
        centered(draw, (cx, bottom - height - 40), f"${value:.2f}B", font(25, True))
        wrapped_center(draw, cx, bottom + 20, label_lines, font(21))
    img.save(market, dpi=(220, 220))

    outlook = BUILD_DIR / "base_outlook.png"
    img = Image.new("RGB", (1400, 680), "white")
    draw = ImageDraw.Draw(img)
    centered(draw, (700, 25), "Base case: operating profit follows a multi-year build", font(34, True))
    left, right, top, bottom = 105, 1340, 125, 555
    max_value, min_value = 750, -100

    def y_for(value):
        return bottom - ((value - min_value) / (max_value - min_value)) * (bottom - top)

    for tick in (-100, 0, 100, 300, 500, 700):
        y = y_for(tick)
        draw.line((left, y, right, y), fill="#D4DCE3", width=2)
        draw.text((18, y - 12), f"{tick}", font=font(19), fill=ink(DARK_GRAY))
    draw.text((18, 82), "$000s", font=font(20, True), fill=ink(DARK_GRAY))
    draw.rectangle((1020, 78, 1050, 98), fill=f"#{NAVY}")
    draw.text((1060, 72), "Net revenue", font=font(20), fill=ink(DARK_GRAY))
    draw.rectangle((1190, 78, 1220, 98), fill=f"#{TEAL}")
    draw.text((1230, 72), "EBITDA", font=font(20), fill=ink(DARK_GRAY))
    years = [2027, 2028, 2029, 2030, 2031]
    revenue = [67, 152, 288, 468, 688]
    ebitda = [-23, -28, -5, 41, 112]
    centers = [250, 490, 730, 970, 1210]
    zero_y = y_for(0)
    for year, rev, ebit, cx in zip(years, revenue, ebitda, centers):
        for value, color, x1, x2 in ((rev, NAVY, cx - 70, cx - 8), (ebit, TEAL, cx + 8, cx + 70)):
            y = y_for(value)
            draw.rectangle((x1, min(y, zero_y), x2, max(y, zero_y)), fill=f"#{color}")
            label_y = y - 28 if value >= 0 else y + 5
            centered(draw, ((x1 + x2) / 2, label_y), f"{value:,}", font(18, True), color)
        centered(draw, (cx, bottom + 15), str(year), font(21), DARK_GRAY)
    img.save(outlook, dpi=(220, 220))

    scenarios = BUILD_DIR / "scenario_ebitda.png"
    img = Image.new("RGB", (1400, 560), "white")
    draw = ImageDraw.Draw(img)
    centered(draw, (700, 24), "Scenario range demonstrates the cost of unproven demand", font(34, True))
    left, right, top, bottom = 270, 1325, 120, 455
    min_value, max_value = -200, 900

    def x_for(value):
        return left + ((value - min_value) / (max_value - min_value)) * (right - left)

    for tick in (-200, 0, 200, 400, 600, 800):
        x = x_for(tick)
        draw.line((x, top, x, bottom), fill="#D4DCE3", width=2)
        centered(draw, (x, bottom + 12), f"{tick}", font(19), DARK_GRAY)
    zero_x = x_for(0)
    names = ["Conservative", "Base", "Upside"]
    values = [-158, 97, 849]
    colors = [RED, TEAL, GOLD]
    ys = [165, 270, 375]
    for name, value, color, cy in zip(names, values, colors, ys):
        draw.text((25, cy - 18), name, font=font(23, True), fill=ink(NAVY))
        value_x = x_for(value)
        draw.rounded_rectangle((min(zero_x, value_x), cy - 25, max(zero_x, value_x), cy + 25), radius=7, fill=f"#{color}")
        label_x = value_x + 20 if value >= 0 else value_x - 20
        box = draw.textbbox((0, 0), f"{value:+,}", font=font(22, True))
        x_text = label_x if value >= 0 else label_x - (box[2] - box[0])
        draw.text((x_text, cy - 15), f"{value:+,}", font=font(22, True), fill=ink(NAVY))
    centered(draw, (800, 520), "Five-year cumulative EBITDA ($000s)", font(21, True), DARK_GRAY)
    img.save(scenarios, dpi=(220, 220))

    return {"market": market, "outlook": outlook, "scenarios": scenarios}


def add_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(6.25))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", caption)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_paragraph_control(p, keep_lines=True)
    cap = doc.add_paragraph()
    cap.style = doc.styles["Caption"]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)
    cap.paragraph_format.space_after = Pt(8)
    set_paragraph_control(cap, keep_lines=True)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    section.different_first_page_header_footer = True

    props = doc.core_properties
    props.title = "Mr Wallaboo Games — Business Prospectus, Feasibility Study & Five-Year Plan"
    props.subject = "Internal planning prospectus"
    props.author = "Mr Wallaboo Games"
    props.keywords = "tabletop games, feasibility, market research, financial plan"
    props.comments = "WORKING v0.1 — illustrative, unaudited planning document"

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(17)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(NAVY)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(TEAL)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(DARK_GRAY)

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.25)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4680, 4680])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.cell(0, 0).text = "MR WALLABOO GAMES"
    table.cell(0, 1).text = "PROSPECTUS · WORKING v0.1"
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_margins(cell, top=0, bottom=50, start=0, end=0)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(NAVY)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), GOLD)
    borders.append(bottom)
    tbl_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INTERNAL PLANNING DOCUMENT  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    add_page_number(p)

    return doc


def add_cover(doc: Document) -> None:
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("BUSINESS PROSPECTUS · FEASIBILITY STUDY · FIVE-YEAR PLAN")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MR WALLABOO\nGAMES")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(7)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("A milestone-gated plan for an original tabletop-game studio and publisher")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    subtitle.paragraph_format.space_after = Pt(20)

    band = doc.add_table(rows=2, cols=2)
    set_table_geometry(band, [2500, 6860])
    band.style = "Table Grid"
    set_table_borders(band, color=GOLD, size=8)
    values = [
        ("RECOMMENDATION", "CONDITIONAL GO"),
        ("CURRENT POSTURE", "Six-month capped validation phase · no manufacturing commitment"),
    ]
    for ridx, (label, value) in enumerate(values):
        for cidx, text in enumerate((label, value)):
            cell = band.cell(ridx, cidx)
            set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
            set_cell_shading(cell, NAVY if cidx == 0 else PALE_GOLD)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10 if cidx == 0 else 12)
            run.font.color.rgb = RGBColor.from_string(WHITE if cidx == 0 else NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    thesis = doc.add_paragraph()
    thesis.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = thesis.add_run("The next investment should buy evidence: one validated game, a cleared brand, a qualified audience, and firm supplier and fulfillment quotes.")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    thesis.paragraph_format.space_before = Pt(12)
    thesis.paragraph_format.space_after = Pt(28)

    meta = doc.add_table(rows=5, cols=2)
    set_table_geometry(meta, [2800, 6560])
    meta_rows = [
        ("Prepared for", "Dan Trezise and Micah Trezise"),
        ("Prepared", "July 17, 2026"),
        ("Planning horizon", "2027–2031"),
        ("Version", "WORKING v0.1"),
        ("Document status", "Internal decision document — not a securities offering"),
    ]
    for ridx, (label, value) in enumerate(meta_rows):
        for cidx, text in enumerate((label, value)):
            cell = meta.cell(ridx, cidx)
            set_cell_margins(cell, top=70, bottom=70, start=90, end=90)
            p = cell.paragraphs[0]
            p.add_run(text)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = cidx == 0
                run.font.color.rgb = RGBColor.from_string(NAVY if cidx == 0 else DARK_GRAY)
    set_table_borders(meta, color=WHITE, size=0)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Illustrative, unaudited planning analysis. See the disclaimer and source appendix.")
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    note.paragraph_format.space_before = Pt(16)
    doc.add_page_break()


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=GOLD, size=8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GOLD)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    add_inline(p, text)
    for run in p.runs:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    set_paragraph_control(p, keep_lines=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body_from_markdown(doc: Document, charts: dict[str, Path]) -> None:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. Executive Summary"))
    lines = lines[start:]
    idx = 0
    body_section_count = 0
    chart_triggers = {
        "6.1 Market size with the right caveat": (charts["market"], "Figure 1. Context-only market measures; category definitions overlap and are not additive. Sources: S01, S03, S05."),
        "11.1 Base-case five-year summary": (charts["outlook"], "Figure 2. Illustrative base case from the linked financial model; values are unaudited planning assumptions."),
        "11.2 Scenario outcomes": (charts["scenarios"], "Figure 3. Five-year cumulative EBITDA range; scenarios are planning cases, not probabilities."),
    }
    bullet_abstract = ensure_list_abstract(doc, "bullet")
    number_abstract = ensure_list_abstract(doc, "number")
    bullet_num_id = new_list_instance(doc, bullet_abstract)
    active_number_num_id = None
    list_mode = None

    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            list_mode = None
            idx += 1
            continue

        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(heading)
            set_paragraph_control(p, keep_next=True, keep_lines=True)
            body_section_count += 1
            list_mode = None
            idx += 1
            continue

        if stripped.startswith("### "):
            heading = stripped[4:].strip()
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(heading)
            set_paragraph_control(p, keep_next=True, keep_lines=True)
            if heading in chart_triggers:
                add_figure(doc, *chart_triggers[heading])
            list_mode = None
            idx += 1
            continue

        if stripped.startswith("#### "):
            heading = stripped[5:].strip()
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(heading)
            set_paragraph_control(p, keep_next=True, keep_lines=True)
            list_mode = None
            idx += 1
            continue

        if stripped.startswith(">"):
            parts = []
            while idx < len(lines) and lines[idx].strip().startswith(">"):
                parts.append(lines[idx].strip()[1:].strip())
                idx += 1
            add_callout(doc, " ".join(parts))
            list_mode = None
            continue

        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            list_mode = None
            continue

        if re.match(r"^-\s+", stripped):
            p = doc.add_paragraph()
            apply_list_numbering(p, bullet_num_id)
            add_inline(p, re.sub(r"^-\s+", "", stripped))
            set_paragraph_control(p, keep_lines=True)
            list_mode = "bullet"
            idx += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            if list_mode != "number":
                active_number_num_id = new_list_instance(doc, number_abstract)
            p = doc.add_paragraph()
            apply_list_numbering(p, active_number_num_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            set_paragraph_control(p, keep_lines=True)
            list_mode = "number"
            idx += 1
            continue

        parts = [stripped]
        idx += 1
        while idx < len(lines):
            nxt = lines[idx].strip()
            if not nxt or nxt.startswith(("#", ">", "|", "- ")) or re.match(r"^\d+\.\s+", nxt):
                break
            parts.append(nxt)
            idx += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(parts))
        set_paragraph_control(p, keep_lines=False, widow=True)
        list_mode = None


def read_source_rows() -> list[list[str]]:
    lines = SOURCE_REGISTER.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("| ID |"))
    rows, _ = parse_table(lines, start)
    return rows


def add_source_appendix(doc: Document) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.add_run("Appendix A — Research Source Register")
    set_paragraph_control(p, keep_next=True, keep_lines=True)
    intro = doc.add_paragraph()
    intro.add_run(
        "Research was current through July 17, 2026. Market totals are contextual, manufacturing figures are budgetary examples rather than quotes, and the name search is not a legal clearance opinion."
    )
    rows = read_source_rows()
    add_table(doc, rows, widths=[650, 1600, 5010, 2100], source_links=True)


def finalize_document(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            set_paragraph_control(paragraph, keep_next=True, keep_lines=True)
    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main() -> None:
    charts = generate_charts()
    doc = setup_document()
    add_cover(doc)
    add_body_from_markdown(doc, charts)
    add_source_appendix(doc)
    finalize_document(doc)
    print(OUTPUT)


if __name__ == "__main__":
    main()
